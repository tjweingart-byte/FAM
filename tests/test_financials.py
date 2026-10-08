"""The finance workbook: every service on one sheet, every day's spend on another.

What is pinned: the recorded spend reaches the ledger on the day it happened;
a usage-billed service reads the ledger rather than a typed number; the plan
the deployment says it bought decides a provider's row; a store that does not
exist is never created by reading it; and only an operator can download it.
"""
from __future__ import annotations

import io
import os
import sys
import time

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

import app as appmod  # noqa: E402
import financials  # noqa: E402
import metering  # noqa: E402

NOW = time.mktime((2026, 10, 7, 12, 0, 0, 0, 0, 0)) - time.timezone


@pytest.fixture
def stores(tmp_path, monkeypatch):
    for var, name in (("METERING_DB", "metering.db"),
                      ("THUMBNAILS_DB", "thumbnails.db"),
                      ("PROVIDER_USAGE_DB", "provider_usage.db")):
        monkeypatch.setenv(var, str(tmp_path / name))
    return tmp_path


def _episode(store, at, *, hit=False):
    u = metering.Usage()
    u.model, u.model_calls = "claude-sonnet-5", 2
    u.input_tokens, u.output_tokens = 6000, 2500
    u.exa_searches, u.exa_cost = 1, 0.015
    u.audio_seconds = 120
    u.cache_hit = hit
    store.record("listener", u, surface="search", minutes=2, at=at)


def _book(data):
    return load_workbook(io.BytesIO(data))


def test_spend_lands_on_the_day_it_happened(stores):
    store = metering.MeterStore()
    _episode(store, NOW)
    _episode(store, NOW - 86400)
    _episode(store, NOW - 86400, hit=True)
    days = financials.daily_ledger(NOW)
    assert days["2026-10-07"]["episodes"] == 1
    assert days["2026-10-06"]["episodes"] == 1
    assert days["2026-10-06"]["hits"] == 1
    claude = metering.price_of(_usage()).claude
    assert days["2026-10-07"]["claude"] == pytest.approx(claude, abs=1e-6)


def _usage():
    u = metering.Usage()
    u.model, u.model_calls = "claude-sonnet-5", 2
    u.input_tokens, u.output_tokens = 6000, 2500
    return u


def test_the_ledger_runs_to_today_with_a_row_for_every_day(stores):
    book = _book(financials.build(NOW))
    led = book["Daily Spend"]
    dates = [led.cell(row=r, column=1).value for r in range(2, led.max_row + 1)]
    assert len(dates) == financials.LEDGER_MIN_DAYS
    assert dates[-1].date().isoformat() == "2026-10-07"


def test_usage_billed_rows_read_the_ledger_never_a_typed_price(stores):
    """A typed number for Claude would be a guess that never moves; the row has
    to follow what was actually spent."""
    cs = _book(financials.build(NOW))["Costs"]
    usage = [r for r in range(10, cs.max_row + 1)
             if str(cs.cell(row=r, column=8).value or "").startswith("Recorded")]
    assert len(usage) >= 3, "Claude, Exa and the GPU are usage-billed"
    for r in usage:
        assert str(cs.cell(row=r, column=6).value).startswith("=SUMIFS('Daily Spend'")


def test_every_service_the_code_calls_has_a_row(stores):
    vendors = {r["vendor"] for r in financials.catalogue()}
    for needed in ("Render", "RunPod", "Anthropic", "Exa", "GNews", "API-Sports",
                   "Finnhub", "Open-Meteo", "Polymarket", "GDELT",
                   "National Weather Service", "Google AI Studio", "Viral Loops",
                   "Apple"):
        assert needed in vendors, needed


def _set(monkeypatch, **changes):
    import dataclasses

    import config

    monkeypatch.setattr(config, "settings",
                        dataclasses.replace(config.settings, **changes))


def test_a_bought_plan_changes_its_row_and_leaves_the_launch_list(stores, monkeypatch):
    _set(monkeypatch, gnews_plan="free")
    free = {r["vendor"]: r for r in financials.catalogue()}["GNews"]
    assert free["price"] == 0 and "licence" in free["status"]
    assert any("GNews" in p["item"] for p in financials.launch_purchases(
        financials.catalogue()))
    _set(monkeypatch, gnews_plan="essential")
    paid = {r["vendor"]: r for r in financials.catalogue()}["GNews"]
    assert paid["status"] == "Paying" and paid["unit"] == "EUR/month"
    assert not any("GNews" in p["item"] for p in financials.launch_purchases(
        financials.catalogue()))


def test_a_pod_is_a_fixed_rental_and_serverless_is_usage(stores, monkeypatch):
    _set(monkeypatch, remote_voice_transport="runpod")
    gpu = [r for r in financials.catalogue() if r["category"] == financials.VOICE][0]
    assert gpu["unit"] == "usage" and gpu["ledger"] == "gpu"
    _set(monkeypatch, remote_voice_transport="http")
    gpu = [r for r in financials.catalogue() if r["category"] == financials.VOICE][0]
    assert gpu["unit"] == "/month" and gpu["price"] > 100
    plan = _book(financials.build(NOW))["Projections"]
    labels = {plan.cell(row=r, column=1).value: r for r in range(1, plan.max_row + 1)}
    episode = plan.cell(row=labels["Cost of one new episode"], column=2).value
    assert "'Daily Spend'!$E:$E" not in episode, "a pod's GPU must not be counted twice"


def test_reading_never_creates_a_store(stores):
    financials.build(NOW)
    for name in ("metering.db", "thumbnails.db", "provider_usage.db"):
        assert not (stores / name).exists(), name


def test_only_an_operator_can_download_it(stores, monkeypatch):
    client = TestClient(appmod.app)
    monkeypatch.setattr(appmod, "ADMIN_TOKEN", "")
    assert client.get("/api/admin/financials.xlsx").status_code == 404
    monkeypatch.setattr(appmod, "ADMIN_TOKEN", "the-real-token")
    assert client.get("/api/admin/financials.xlsx",
                      headers={"X-Admin-Token": "guess"}).status_code == 404
    res = client.get("/api/admin/financials.xlsx",
                     headers={"X-Admin-Token": "the-real-token"})
    assert res.status_code == 200
    assert "FAM_Financials_" in res.headers["content-disposition"]
    assert _book(res.content).sheetnames == [
        "Costs", "Marketing & Materials", "Projections", "Daily Spend"]


def test_the_plan_is_driven_by_listeners_with_sports_apart(stores):
    """Cost follows new episodes, and sports misses the cache more and spends
    API-Sports requests - so it has its own rows, not a blended growth rate."""
    plan = _book(financials.build(NOW))["Projections"]
    labels = {plan.cell(row=r, column=1).value: r for r in range(1, plan.max_row + 1)}
    for needed in ("Listeners", "  of which sports",
                   "New episodes per search, sports", "Sports data (API-Sports)",
                   "Marketing & materials", "Sports share of running costs",
                   "Total costs"):
        assert needed in labels, needed
    episode = plan.cell(row=labels["Cost of one new episode"], column=2).value
    assert "'Daily Spend'" in episode, "the forecast must learn from recorded spend"
    api = plan.cell(row=labels["Sports data (API-Sports)"], column=4).value
    assert "7500" not in api, "plan limits live in labelled rate cells"
    rate = labels["API-Sports Pro: requests per day"]
    assert plan.cell(row=rate, column=2).value == 7500 and f"$B${rate}" in api


def test_marketing_and_materials_feed_the_projection(stores):
    """Every month of the plan carries that month's marketing budget, and the
    launch month carries the one-off spend too."""
    book = _book(financials.build(NOW))
    mkt, plan = book["Marketing & Materials"], book["Projections"]
    lines = [mkt.cell(row=r, column=1).value for r in range(1, mkt.max_row + 1)]
    assert "Marketing" in lines and "Materials" in lines
    assert len(financials.MARKETING) == sum(
        1 for line in financials.MARKETING if line[1] in lines)
    labels = {plan.cell(row=r, column=1).value: r for r in range(1, plan.max_row + 1)}
    row = labels["Marketing & materials"]
    assert plan.cell(row=row, column=4).value.startswith("='Marketing & Materials'!E")
    first = mkt.cell(row=7, column=5).value
    assert "$B$3" in first and "$D7" in first, "before/after launch and the one-off"


def test_the_google_sheet_gets_the_same_days_as_the_workbook(stores, monkeypatch):
    """The Apps Script replaces the sheet's Daily Spend rows with this JSON, so
    it must be the workbook's own rows, admin-only like the download."""
    store = metering.MeterStore()
    _episode(store, time.time())
    client = TestClient(appmod.app)
    monkeypatch.setattr(appmod, "ADMIN_TOKEN", "the-real-token")
    assert client.get("/api/admin/financials/daily.json").status_code == 404
    res = client.get("/api/admin/financials/daily.json",
                     headers={"X-Admin-Token": "the-real-token"})
    assert res.status_code == 200
    body = res.json()
    assert body["rows"][-1][0] == body["as_of"]
    assert body["rows"][-1][1] == 1 and body["rows"][-1][2] > 0
    assert len(body["rows"]) == financials.LEDGER_MIN_DAYS
    script = open(os.path.join(os.path.dirname(os.path.dirname(
        os.path.abspath(__file__))), "tools", "financials_apps_script.gs")).read()
    assert "/api/admin/financials/daily.json" in script
    assert "'Daily Spend'" in script and "FAM_ADMIN_TOKEN" in script


def test_ledger_ranges_are_whole_columns_and_any_launch_day_counts(stores):
    """The Google Sheet adds a row a day, so a range fixed at today's last row
    would leave new days out; and 15 Jan as the launch month means January."""
    book = _book(financials.build(NOW))
    cs = book["Costs"]
    usage = [cs.cell(row=r, column=6).value for r in range(10, cs.max_row + 1)
             if str(cs.cell(row=r, column=6).value or "").startswith("=SUMIFS")]
    assert usage and all("'Daily Spend'!$A:$A" in f for f in usage)
    plan = book["Projections"]
    listeners = [plan.cell(row=r, column=4).value for r in range(1, plan.max_row + 1)
                 if plan.cell(row=r, column=1).value == "Listeners"][0]
    assert "DATE(YEAR($B$4),MONTH($B$4),1)" in listeners
    assert book["Marketing & Materials"]["B3"].value.startswith("=DATE(YEAR(")


def test_a_photo_check_is_spend_but_not_an_episode(stores):
    """§225's photo check records Claude spend in the same ledger. It must show
    in the day's Claude cost and never in the episode count, which divides it."""
    store = metering.MeterStore()
    _episode(store, NOW)
    u = metering.Usage()
    u.model, u.model_calls = "claude-haiku-5-5", 1
    u.input_tokens, u.output_tokens = 400, 50
    store.record("listener", u, surface="photo_check", at=NOW)
    day = financials.daily_ledger(NOW)["2026-10-07"]
    assert day["episodes"] == 1
    assert day["claude"] == pytest.approx(
        metering.price_of(_usage()).claude + metering.price_of(u).claude, abs=1e-6)
